from pptx import Presentation

def parse_pptx(path):

    prs = Presentation(path)

    texts = []

    for slide in prs.slides:

        for shape in slide.shapes:

            if hasattr(
                shape,
                "text"
            ):
                texts.append(
                    shape.text
                )

    return "\n".join(texts)